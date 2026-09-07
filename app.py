import os
from io import BytesIO

import pandas as pd
import streamlit as st

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# ============================================================
# MDC SOLUTION - VERSION 1 (REAL SINGLE-RACK DATA)
# Data source: MDC_Master_V1.xlsx (same folder as app.py)
#
# Single Rack:
#   Configuration 1-4 = real data from supplied MDC BOQ
#
# Multirack:
#   Configuration 1-9 = XXX placeholders for future update
# ============================================================

st.set_page_config(
    page_title="MDC Solution",
    page_icon="🏢",
    layout="wide",
    initial_sidebar_state="expanded",
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
MASTER_FILE = os.path.join(BASE_DIR, "MDC_Master_V1.xlsx")

DEMO_INTERNAL_PASSWORD = "MDC@123"  # Change before production.

# ------------------------------------------------------------
# Eaton / MDC UI styling
# ------------------------------------------------------------

st.markdown(
    """
    <style>
    :root {
        --mdc-blue: #0167C9;
        --mdc-dark-blue: #004B91;
        --mdc-light-blue: #EAF3FC;
    }
    .stApp { background-color: #ffffff; }
    h1, h2, h3 { color: var(--mdc-dark-blue); }
    .stButton > button {
        background-color: var(--mdc-blue);
        color: white;
        border: none;
        border-radius: 6px;
        font-weight: 600;
    }
    .stButton > button:hover {
        background-color: var(--mdc-dark-blue);
        color: white;
    }
    input:focus, textarea:focus {
        border-color: var(--mdc-blue) !important;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# ------------------------------------------------------------
# Load master data
# ------------------------------------------------------------

@st.cache_data
def load_master():
    configs = pd.read_excel(MASTER_FILE, sheet_name="Configurations")
    components = pd.read_excel(MASTER_FILE, sheet_name="Components")
    accessories = pd.read_excel(MASTER_FILE, sheet_name="Accessories")
    pdus = pd.read_excel(MASTER_FILE, sheet_name="PDUs")
    return configs, components, accessories, pdus

configs_df, components_df, accessories_df, pdus_df = load_master()

# ------------------------------------------------------------
# Session state defaults
# ------------------------------------------------------------

defaults = {
    "mode": "Sales",
    "authenticated": False,
    "customer_name": "",
    "customer_place": "",
    "problem": "",
    "solution": "",
    "mdc_type": "Single Rack",
    "configuration": "Configuration 1",
    "accessory_qty": {},
    "pdu_qty": {},
    "margin_pct": 20.0,
    "freight": 0.0,
    "installation": 0.0,
    "warranty_pct": 0.0,
}

for key, value in defaults.items():
    if key not in st.session_state:
        st.session_state[key] = value

# ------------------------------------------------------------
# Helper functions
# ------------------------------------------------------------

def money(value):
    try:
        return f"₹ {float(value):,.2f}"
    except Exception:
        return "₹ 0.00"

def price_box(label, value):
    st.markdown(
        f"""
        <div style="padding:4px 0 12px 0; min-height:82px; overflow:visible;">
            <div style="font-size:16px; color:#4b5563; margin-bottom:7px;">
                {label}
            </div>
            <div style="font-size:30px; font-weight:600; color:#30333d; white-space:nowrap; overflow:visible;">
                {money(value)}
            </div>
        </div>
        """,
        unsafe_allow_html=True,
    )

def internal_password():
    try:
        return st.secrets["MDC_INTERNAL_PASSWORD"]
    except Exception:
        return DEMO_INTERNAL_PASSWORD

def selected_config_record():
    match = configs_df[
        (configs_df["MDC Type"] == st.session_state.mdc_type)
        & (configs_df["Configuration"] == st.session_state.configuration)
    ]
    return match.iloc[0] if not match.empty else None

def selected_components():
    return components_df[
        (components_df["MDC Type"] == st.session_state.mdc_type)
        & (components_df["Configuration"] == st.session_state.configuration)
    ].copy()

# ------------------------------------------------------------
# Build BOM
# ------------------------------------------------------------

def build_bom():
    rows = []
    # Configuration BOM
    for _, r in selected_components().iterrows():
        cost = r["Unit Cost"]
        qty = float(r["Quantity"])
        rows.append({
            "S.No.": len(rows) + 1,
            "Component Type": "Base (Configuration)",
            "Part Code": r["Part Code"] if pd.notna(r["Part Code"]) else "",
            "Description": r["Description"],
            "Quantity": qty,
            "UOM": r["UOM"],
            "Unit Cost": cost,
            "Total Cost": cost * qty if pd.notna(cost) else None,
            "Source": "Configuration",
        })
    # Accessories
    for _, r in accessories_df.iterrows():
        part = str(r["Part Code"])
        qty = float(st.session_state.accessory_qty.get(part, 0))
        if qty > 0:
            cost = r["Unit Cost"]
            rows.append({
                "S.No.": len(rows) + 1,
                "Component Type": "Optional Accessory",
                "Part Code": part if part.strip() else "",
                "Description": r["Description"],
                "Quantity": qty,
                "UOM": r["UOM"],
                "Unit Cost": cost,
                "Total Cost": cost * qty if pd.notna(cost) else None,
                "Source": "Optional Accessory",
            })
    # PDUs
    for _, r in pdus_df.iterrows():
        part = str(r["Part Code"])
        qty = float(st.session_state.pdu_qty.get(part, 0))
        if qty > 0:
            cost = r["Unit Cost"]
            desc = f'{r["Description"]} | Type: {r["Type"]} | C13: {r["C13"]} | C19: {r["C19"]}'
            rows.append({
                "S.No.": len(rows) + 1,
                "Component Type": "PDU",
                "Part Code": part if part.strip() else "",
                "Description": desc,
                "Quantity": qty,
                "UOM": r["UOM"],
                "Unit Cost": cost,
                "Total Cost": cost * qty if pd.notna(cost) else None,
                "Source": "PDU",
            })
    return pd.DataFrame(rows)

# ------------------------------------------------------------
# Cost summary
# ------------------------------------------------------------

def cost_summary(bom):
    cfg = selected_config_record()
    base_cost = float(cfg["Base Cost"]) if cfg is not None and pd.notna(cfg["Base Cost"]) else 0.0
    optional_cost = bom.loc[bom["Source"]=="Optional Accessory","Total Cost"].fillna(0).sum() if not bom.empty else 0.0
    pdu_cost = bom.loc[bom["Source"]=="PDU","Total Cost"].fillna(0).sum() if not bom.empty else 0.0
    total_cost = base_cost + optional_cost + pdu_cost
    return base_cost, optional_cost, pdu_cost, total_cost

# ------------------------------------------------------------
# Selling price calculation
# ------------------------------------------------------------

def add_selling_prices(bom, total_cost, margin_pct, freight, installation):
    result = bom.copy()
    margin_price = total_cost / (1 - margin_pct/100) if margin_pct < 100 else 0
    final_selling_price = margin_price + freight + installation
    known_cost_total = result["Total Cost"].fillna(0).sum() if not result.empty else 0
    if known_cost_total > 0:
        result["Total Price"] = result["Total Cost"].fillna(0) / known_cost_total * final_selling_price
        result["Unit Price"] = result["Total Price"] / result["Quantity"]
    else:
        result["Total Price"] = pd.NA
        result["Unit Price"] = pd.NA
    return result, margin_price, final_selling_price

# ------------------------------------------------------------
# Customer table
# ------------------------------------------------------------

def customer_table():
    return pd.DataFrame([
        ["Customer Name", st.session_state.customer_name],
        ["Customer Place", st.session_state.customer_place],
        ["Problem Description", st.session_state.problem],
        ["Solution", st.session_state.solution],
        ["MDC Type", st.session_state.mdc_type],
        ["Configuration", st.session_state.configuration],
    ], columns=["Field","Value"])

# ------------------------------------------------------------
# Excel helper
# ------------------------------------------------------------

def write_dataframe(ws, dataframe, start_row, start_col=1, header=True):
    dataframe = dataframe.copy().where(pd.notna(dataframe), None)
    current_row = start_row
    if header
